"""Spreadsheet DPR ingest — schema inference, not NLP (spec §4.1).

* find the header row (the row whose cells best match known column synonyms — DPRs have
  title rows above the header)
* map each header to a canonical field with rapidfuzz ("Desc of Work" → description,
  "Qty Done" → quantity, "Nos" → uom, "Manpwr" → manpower …)
* one row = one candidate event. The description/remarks cells are read with the rule
  vocabulary; the mapped quantity/uom/location/manpower columns override it (their
  spans point at the cell inside the row line we build as `raw_text`).
"""
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz import fuzz, process

from . import extract_rules
from .schema import ExtractedEvent

SYNONYMS = {
    "description": ["description", "desc of work", "description of work", "activity", "work done", "item", "work description", "particulars", "desc"],
    "location": ["location", "area", "locn", "area / locn", "area/location", "zone", "work front", "unit"],
    "quantity": ["qty", "qty done", "quantity", "quantity done", "today qty", "progress qty", "done"],
    "uom": ["uom", "unit of measure", "nos", "units", "u/m"],
    "manpower": ["manpower", "manpwr", "mp", "labour", "workers", "man power", "manpower deployed"],
    "remarks": ["remarks", "status", "comment", "comments", "note"],
    "date": ["date", "dt", "work date"],
    "sl": ["sl", "sl no", "s no", "sr", "sr no", "#"],
}
_choices = {syn: field for field, syns in SYNONYMS.items() for syn in syns}


def map_headers(headers: list[str], min_score: int = 80) -> dict[int, dict]:
    """→ {col_index: {field, header, score}} — the fuzzy header mapping, shown to the user."""
    out, taken = {}, set()
    for i, h in enumerate(headers):
        h0 = str(h or "").strip().lower()
        if not h0:
            continue
        best = process.extractOne(h0, list(_choices), scorer=fuzz.token_set_ratio)
        if best and best[1] >= min_score:
            field = _choices[best[0]]
            if field not in taken:
                out[i] = {"field": field, "header": str(h), "score": round(best[1])}
                taken.add(field)
    return out


def find_header_row(rows: list[list]) -> tuple[int, dict]:
    best = (-1, {})
    for i, r in enumerate(rows[:12]):
        m = map_headers(r)
        if len(m) > len(best[1]):
            best = (i, m)
    return best


def _find_date(rows) -> date | None:
    for r in rows[:6]:
        for c in r:
            if isinstance(c, datetime):
                return c.date()
            m = re.search(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", str(c or ""))
            if m:
                y = int(m.group(3))
                return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
    return None


def read(path: Path, fallback_date: date) -> dict:
    """→ {raw_text, reporting_date, mapping, events: [ExtractedEvent]}"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi, mapping = find_header_row(rows)
    rdate = _find_date(rows) or fallback_date
    field_col = {v["field"]: k for k, v in mapping.items()}
    lines, events, pos = [], [], 0
    for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(c not in (None, "") for c in r):
            continue
        cells = {f: r[c] if c < len(r) else None for f, c in field_col.items()}
        if not cells.get("description"):
            continue
        # build a readable row line; record where each cell lands so spans stay exact
        parts, spans = [], {}
        cursor = pos + len(f"Row {ri}: ")
        for f in ("description", "location", "quantity", "uom", "manpower", "remarks"):
            v = cells.get(f)
            if v in (None, ""):
                continue
            txt = f"{v:g}" if isinstance(v, float) else str(v)
            spans[f] = (cursor, cursor + len(txt))
            parts.append(txt)
            cursor += len(txt) + 3
        line = f"Row {ri}: " + " | ".join(parts)
        lines.append(line)
        # read the description + remarks with the rule vocabulary (no LLM)
        full = "\n".join(lines)
        d0 = spans["description"][0]
        sub = extract_rules._extract_one(full, d0, pos + len(line), rdate, "en") or \
            ExtractedEvent(sentence=line, sentence_span=(pos, pos + len(line)), extraction_conf=0.6)
        sub.sentence, sub.sentence_span = line, (pos, pos + len(line))
        sub.extractor = "spreadsheet"
        # mapped columns win over free-text reading
        if cells.get("location") not in (None, ""):
            sub.location_text = str(cells["location"])
            sub.spans["location_text"] = spans["location"]
        if isinstance(cells.get("quantity"), (int, float)):
            sub.quantity = float(cells["quantity"])
            sub.spans["quantity"] = spans["quantity"]
        if cells.get("uom"):
            sub.uom = str(cells["uom"]).lower()
            sub.spans["uom"] = spans["uom"]
        if isinstance(cells.get("manpower"), (int, float)):
            sub.manpower_total = int(cells["manpower"])
            sub.spans["manpower_total"] = spans["manpower"]
        sub.date_conf = 0.9 if rdate != fallback_date else 0.85
        sub.extraction_conf = round(min(0.95, max(sub.extraction_conf, 0.8) + 0.05), 3)   # structured columns are reliable
        events.append(sub)
        pos += len(line) + 1
    return {"raw_text": "\n".join(lines), "reporting_date": rdate, "mapping": mapping, "header_row": hi + 1, "events": events}


def read_pdf_text(path: Path) -> str:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


HEADING = re.compile(r"^\s*(?:[A-Z0-9][A-Z0-9 &/\-–]{3,}|.{3,40}:)\s*$")


def dpr_sections(text: str):
    """→ section_of(pos) for DPR text split by area/system headings (ALL CAPS lines)."""
    marks = []
    pos = 0
    for line in text.splitlines(keepends=True):
        if HEADING.match(line.strip()) and len(line.strip()) < 60 and not re.search(r"\d{1,2}[-/]\w{3}[-/]\d{4}", line):
            marks.append((pos, line.strip().rstrip(":")))
        pos += len(line)

    def section_of(p):
        cur = None
        for s, h in marks:
            if s <= p:
                cur = h
        return cur
    return section_of


CHAT_LINE = re.compile(r"^\[(\d{1,2})/(\d{1,2})/(\d{2,4}),\s*(\d{1,2}):(\d{2})\s*(AM|PM)?\]\s*([^:]+):\s*(.+)$", re.I)


def parse_chat_export(text: str):
    """WhatsApp '[dd/mm/yy, h:mm AM] Name: message' lines → [(datetime, name, message)]."""
    out = []
    for line in text.splitlines():
        m = CHAT_LINE.match(line.strip())
        if not m:
            continue
        d, mo, y, h, mi, ap, name, msg = m.groups()
        y = int(y)
        h = int(h)
        if ap and ap.upper() == "PM" and h < 12:
            h += 12
        if ap and ap.upper() == "AM" and h == 12:
            h = 0
        out.append((datetime(y + 2000 if y < 100 else y, int(mo), int(d), h, int(mi)), name.strip(), msg.strip()))
    return out
